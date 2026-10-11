"""Editable PowerPoint from the verified slide data; photos use native PPTX crops."""
from pathlib import Path
import json
from pptx import Presentation
from pptx.util import Inches,Pt
from pptx.dml.color import RGBColor
from PIL import Image
ROOT=Path(__file__).resolve().parents[1]
slides=json.loads((ROOT/'05_praesentation/deck.json').read_text())
assets={x['id']:x for x in json.loads((ROOT/'01_recherche/asset_manifest.json').read_text())}
import sys
sys.path.insert(0,str(ROOT/'werkzeuge'))
from content import SOURCES
prs=Presentation();prs.slide_width=Inches(13.333333);prs.slide_height=Inches(7.5)
prs.core_properties.author='Johannes Fink';prs.core_properties.title='OMA / LMN — Seattle Central Library, Seattle, 2004';prs.core_properties.subject='Architektonische Objektanalyse'
INK='16221B';LIGHT='F4F6ED';LIME='D9EF65';MUTED='CCD1C2'
def txt(sl,text,x,y,w,h,size=20,color=LIGHT,font='Arial',bold=False,url=None):
 box=sl.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h));tf=box.text_frame;tf.clear();tf.word_wrap=True
 tf.margin_left=0;tf.margin_right=0;tf.margin_top=0;tf.margin_bottom=0
 for i,line in enumerate(text.split('\n')):
  p=tf.paragraphs[0] if i==0 else tf.add_paragraph();p.text=line;p.font.name=font;p.font.size=Pt(size);p.font.bold=bold;p.font.color.rgb=RGBColor.from_string(color)
  if url:
   for r in p.runs:r.hyperlink.address=url
 return box

def picture(sl,key,x,y,w,h,cover=False):
 a=assets[key];path=ROOT/a['path']
 if path.suffix=='.svg':path=path.with_suffix('.png')
 iw,ih=Image.open(path).size;ar=iw/ih;boxar=w/h
 if cover:
  pic=sl.shapes.add_picture(str(path),Inches(x),Inches(y),width=Inches(w),height=Inches(h))
  if ar>boxar:crop=(1-boxar/ar)/2;pic.crop_left=crop;pic.crop_right=crop
  else:crop=(1-ar/boxar)/2;pic.crop_top=crop;pic.crop_bottom=crop
 else:
  scale=min(w/iw,h/ih);pw,ph=iw*scale,ih*scale
  pic=sl.shapes.add_picture(str(path),Inches(x+(w-pw)/2),Inches(y+(h-ph)/2),width=Inches(pw),height=Inches(ph))
 pic._element.nvPicPr.cNvPr.set('descr',a['title']+'; '+a['author']+'; '+a['license'])
 txt(sl,key+' / '+a['title'],x,y+h+.08,w,.2,8.2,MUTED)
 txt(sl,a['author']+' · '+a['license'],x,y+h+.25,w,.26,7.2,MUTED,url=a['url'])

for s in slides:
 sl=prs.slides.add_slide(prs.slide_layouts[6]);sl.background.fill.solid();sl.background.fill.fore_color.rgb=RGBColor.from_string(INK)
 if s.get('layout')=='cover':
  txt(sl,s['kicker'],.5,.42,6,.3,10,LIME,font='Courier New')
  txt(sl,s['title'],.5,1,6,2.1,54,font='Georgia')
  txt(sl,s['lines'][0],.5,3.4,5.3,.8,17)
  picture(sl,'P02',.5,4.5,2.1,1.72)
  picture(sl,'F01',6.8,.42,6,5.85,cover=True)
 else:
  txt(sl,s['kicker'],.5,.42,12,.3,10,LIME,font='Courier New')
  title_size=34 if len(s['title'])<50 else 30
  txt(sl,s['title'],.5,.85,12.3,.75,title_size,font='Georgia')
  layout=s.get('layout')
  if layout=='sources':
   for i,(code,title,url,desc) in enumerate(SOURCES):
    col=i%2;row=i//2;x=.5+col*6.2;y=1.8+row*.51
    txt(sl,'['+code+'] '+title,x,y,6,.28,9.2,bold=True)
    txt(sl,url.replace('https://',''),x,y+.23,6,.25,7,LIME,url=url)
  elif layout in ['credits','credits_plans']:
   used=sorted(set(k for d in slides for k in d['images'] if (k.startswith('F') if layout=='credits' else k.startswith('P'))))
   for i,key in enumerate(used):
    a=assets[key];col=i%3;row=i//3;x=.5+col*4.13;y=1.8+row*.55
    txt(sl,key+' / '+a['title'],x,y,3.93,.23,8.2,bold=True)
    txt(sl,a['author'],x,y+.22,3.93,.18,7.4,MUTED)
    txt(sl,a['license'],x,y+.38,3.93,.18,7.3,LIME,url=a['url'])
   if layout=='credits_plans':txt(sl,'A01–A05: Johannes Fink · eigene Analyseskizzen · CC BY 4.0 · ohne Maßstab',.5,6.55,12,.24,8.5,MUTED)
  else:
   lines=s['lines'];n=len(lines);w=(12.33-.25*(n-1))/n
   for i,t in enumerate(lines):txt(sl,t,.5+i*(w+.25),1.66,w,.7,12)
   images=s['images'];n=len(images);w=(12.33-.17*(n-1))/n
   for i,key in enumerate(images):picture(sl,key,.5+i*(w+.17),2.5,w,3.72,cover=key.startswith('F'))
 txt(sl,'Quellen: '+s['refs'],.5,7.04,9.8,.23,7.4,'B7C0AA',font='Courier New')
 txt(sl,'Johannes Fink / '+str(s['number']).zfill(2),10.55,7.04,2.4,.23,8,'B7C0AA',font='Courier New')
 sl.notes_slide.notes_text_frame.text=f"(ca. {s['seconds']} s)\n\n{s['notes']}\n\nQuellen: {s['refs']}\nMaterialien: {', '.join(s['images'])}"
prs.save(ROOT/'05_praesentation/OMA_Seattle_Central_Library_Praesentation_Fink.pptx')
print('PowerPoint written: 16 slides, editable text, native photo crops, speaker notes, hyperlinks.')
